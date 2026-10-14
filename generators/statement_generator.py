import sys
from pathlib import Path
from copy import copy

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Border, Side

def resource_path(relative_path):
    """Get the correct path for development and packaged applications."""
    if getattr(sys, "frozen", False):
        base_path = Path(sys._MEIPASS)
    else:
        base_path = Path(__file__).resolve().parent.parent

    return base_path / relative_path
def find_label(ws, labels):
    """Find a label in the worksheet, ignoring letter case."""
    labels = {label.strip().lower() for label in labels}

    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue

            if isinstance(cell.value, str):
                value = cell.value.strip().lower().rstrip(":")
                if value in labels:
                    return cell

    return None


def fill_next_cell(ws, labels, value):
    """Fill the cell immediately to the right of a label."""
    label_cell = find_label(ws, labels)

    if label_cell is None or value is None:
        return

    target = ws.cell(label_cell.row, label_cell.column + 1)

    # If the target belongs to a merged range, use its top-left cell.
    for merged_range in ws.merged_cells.ranges:
        if target.coordinate in merged_range:
            target = ws.cell(
                merged_range.min_row,
                merged_range.min_col
            )
            break

    target.value = value




def find_module_table(ws):
    """Locate the module table using its known column headings."""

    aliases = {
        "semester": "semester",
        "module code": "code",
        "module": "code",
        "code": "code",
        "credit": "credits",
        "credits": "credits",
        "module name": "name",
        "ca": "ca",
        "se": "exam",
        "tot": "total",
        "total": "total",
        "grade": "grade",
        "result": "result",
    }

    for row_number in range(1, ws.max_row + 1):
        headings = {}

        for column in range(1, ws.max_column + 1):
            value = ws.cell(row_number, column).value

            if not isinstance(value, str):
                continue

            key = " ".join(
                value.strip().lower().replace(":", "").split()
            )

            if key in aliases:
                headings[aliases[key]] = column

        # The template's module table uses these headings.
        if (
            "semester" in headings
            and "credits" in headings
            and "name" in headings
            and "ca" in headings
            and "exam" in headings
            and "total" in headings
            and "grade" in headings
            and "result" in headings
        ):
            # Column B is the module-code column in this template.
            headings["code"] = 2
            return row_number, headings

    raise ValueError(
        "Could not locate the module table. "
        "Expected the module headings on the template worksheet."
    )




def copy_row_style(ws, source_row, target_row):
    """Copy formatting from the template's first module row."""
    for column in range(1, ws.max_column + 1):
        source = ws.cell(source_row, column)
        target = ws.cell(target_row, column)

        if isinstance(source, MergedCell) or isinstance(target, MergedCell):
            continue

        if source.has_style:
            target._style = copy(source._style)

        if source.number_format:
            target.number_format = source.number_format

        if source.alignment:
            target.alignment = copy(source.alignment)

        if source.protection:
            target.protection = copy(source.protection)


def generate_student_statement(
    student,
    template_path,
    output_folder,
    programme="",
    examination_period="",
    centre="",
    publication_date="",
    issued_on="",
    semester="I",
):
    """Generate one Excel statement for one student."""
    template_path = Path(template_path)
    output_folder = Path(output_folder)

    if not template_path.is_file():
        raise FileNotFoundError(
            f"Statement template not found: {template_path}"
        )

    registration_number = str(student.registration_number).strip()

    if not registration_number:
        raise ValueError("A student has no registration number.")

    output_folder.mkdir(parents=True, exist_ok=True)

    # Load a fresh copy of the template for each student.
    workbook = load_workbook(template_path)
    ws = workbook.active
    
    print("\n--- TEMPLATE CELL DIAGNOSTIC ---")

    for row in ws.iter_rows():
        cells = [
            f"{cell.coordinate}={cell.value!r}"
            for cell in row
            if cell.value is not None
        ]

        if cells:
            print(" | ".join(cells))

    print("\n--- MERGED CELL RANGES ---")
    print(list(ws.merged_cells.ranges))
    print("--- END DIAGNOSTIC ---\n")



    # Fill the exact cells used by the statement template.
    ws["B6"] = student.name
    ws["J6"] = student.registration_number
    ws["B7"] = programme
    ws["J7"] = examination_period
    ws["B8"] = centre
    ws["F8"] = student.date_of_entry
    if student.date_of_entry:
        ws["F8"].number_format = "mmm-yyyy"
    ws["J8"] = publication_date


    # Locate the module table.
    header_row, columns = find_module_table(ws)
    first_data_row = header_row + 1

    # This template originally contains six module rows.
    template_module_rows = 6
    summary_row = first_data_row + template_module_rows

    # Remove existing merged ranges from the module rows so they
    # can be resized safely.
    for row_number in range(first_data_row, summary_row):
        cell_range = f"D{row_number}:F{row_number}"
        if cell_range in {str(r) for r in ws.merged_cells.ranges}:
            ws.unmerge_cells(cell_range)

    # Unmerge the summary label before inserting or deleting rows.
    summary_merge = f"A{summary_row}:B{summary_row}"
    if summary_merge in {str(r) for r in ws.merged_cells.ranges}:
        ws.unmerge_cells(summary_merge)

    module_count = len(student.modules)
    difference = module_count - template_module_rows

    if difference > 0:
        # Add extra rows immediately before the summary section.
        ws.insert_rows(summary_row, amount=difference)
    elif difference < 0:
        # Remove unused module rows to avoid blank rows.
        ws.delete_rows(
            first_data_row + module_count,
            amount=-difference
        )

    # Restore the merged module-name cells for the actual module rows.
    for row_number in range(first_data_row, first_data_row + module_count):
        ws.merge_cells(f"D{row_number}:F{row_number}")

    # Restore the summary label merge at its new location.
    new_summary_row = summary_row + difference
    ws.merge_cells(f"A{new_summary_row}:B{new_summary_row}")

    # Copy the first module row's formatting to additional module rows.
    for row_number in range(first_data_row + 1, first_data_row + module_count):
        copy_row_style(ws, first_data_row, row_number)


    # Write the module results.

    # Write the student's module results.
    for index, module in enumerate(student.modules):
        row_number = first_data_row + index

        if index > 0:
            copy_row_style(ws, first_data_row, row_number)

        result = ""
        grade = str(module.grade or "").strip().upper()

        if grade:
            result = (
                "FAIL" if grade in {"F", "FAIL"} else "PASS"
            )

        values = {
            "semester": semester,
            "code": module.code,
            "credits": module.credits,
            "name": module.name,
            "ca": module.ca,
            "exam": module.exam,
            "total": module.total,
            "grade": module.grade,
            "result": result,
        }

        for field, value in values.items():
            if field in columns:
                ws.cell(row_number, columns[field]).value = value


    # Fill GPA and related summary fields where matching labels exist.
    fill_next_cell(ws, ["cgpa", "cumulative gpa"], student.cgpa)
    fill_next_cell(
        ws, ["course attendance", "attendance"], student.course_attendance
    )
    fill_next_cell(ws, ["remarks", "remark"], student.remarks)


    # Fill GPA and issued-on date below the module table.
    gpa_row = first_data_row + module_count
    ws.cell(gpa_row, 11).value = student.semester_gpa

    issued_row = gpa_row + 2
    ws.cell(issued_row, 2).value = issued_on

    
    # Set the height and bottom borders of dynamically added module rows.
    if module_count > template_module_rows:
        bottom_side = Side(style="thin")

        for row_number in range(
            first_data_row + template_module_rows,
            first_data_row + module_count,
        ):
            ws.row_dimensions[row_number].height = 45.75

            for column in (5, 6):  # Columns E and F
                cell = ws.cell(row=row_number, column=column)
                cell.border = Border(
                    left=copy(cell.border.left),
                    right=copy(cell.border.right),
                    top=copy(cell.border.top),
                    bottom=bottom_side,
                    diagonal=copy(cell.border.diagonal),
                    diagonal_direction=cell.border.diagonal_direction,
                    diagonalUp=cell.border.diagonalUp,
                    diagonalDown=cell.border.diagonalDown,
                    outline=cell.border.outline,
                    vertical=copy(cell.border.vertical),
                    horizontal=copy(cell.border.horizontal),
                )

    # Find the COE row after the module table has been resized.
    coe_row = None
    for row in ws.iter_rows():
        for cell in row:
            if (
                isinstance(cell.value, str)
                and cell.value.strip().upper() == "COE"
            ):
                coe_row = cell.row
                break
        if coe_row is not None:
            break

    # Print the complete statement through the COE section.
    if coe_row is not None:
        ws.print_area = f"A1:K{max(coe_row, issued_row)}"


    output_path = output_folder / f"{registration_number}.xlsx"

    try:
        workbook.save(output_path)
    finally:
        workbook.close()

    return output_path
